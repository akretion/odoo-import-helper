from typing import dataclass_transform
from decorator import append
from openpyxl.workbook import child
from odoo import api, fields, models
from tempfile import NamedTemporaryFile
import base64

from odoo.fields import Command

import openpyxl as opx

import logging
from stdnum.iban import is_valid as iban_is_valid

logger = logging.getLogger(__name__)

LIST_FR = ["FR", "France", "Martinique - France", "REUNION", "Paris"]
LIST_COL_POP = [
    "invoice_name",
    "invoice_title_code",
    "invoice_street",
    "invoice_street2",
    "invoice_city",
    "invoice_zip",
    "invoice_phone",
    "invoice_mobile",
    "invoice_lang",
    "delivery_name",
    "delivery_title_code",
    "delivery_street",
    "delivery_street2",
    "delivery_city",
    "delivery_zip",
    "delivery_phone",
    "delivery_mobile",
    "delivery_lang",
    "attributes",
    "ref_supplier",
    "ref_template",
    "Colonnes:",
]


class ImportHelpergeneric(models.TransientModel):
    _name = "import.helper.generic"
    _description = "Import helper generic for importing information with template"

    file_import = fields.Binary(string="File to import")
    update_on = fields.Boolean(string="Import or update")

    def speedy_partner_categori_id(self):
        categ_id = {}
        categs = self.env["res.partner.category"].search([])
        for c in categs:
            categ_id[c.name] = c.id
        return categ_id

    def speedy_partner_id(self):
        partner_list = {}
        partner_ids = self.env["res.partner"].search([])
        for p in partner_ids:
            partner_list[p.ref] = p.id
        return partner_list

    def check_vals_partner(self, vals):
        if "category_id" in vals:
            categ_ids = self.speedy_partner_categori_id()
            list_categ = []
            for ctname in vals["category_id"].split("/"):
                if ctname and vals["category_id"] in categ_ids:
                    list_categ.append(categ_ids[ctname])
                vals["category_id"] = [Command.set(list_categ)]
        if "invoice_name" in vals:
            child_val = {"type": "invoice"}
            vals["child_ids"] = []
            for v in vals:
                if "invoice_" in v:
                    child_val[v.replace("invoice_", "")] = vals[v]
            vals["child_ids"].append((0, 0, child_val))
        if "delivery_name" in vals:
            child_val = {"type": "delivery"}
            for v in vals:
                if "delivery_" in v:
                    child_val[v.replace("delivery_", "")] = vals[v]
            if "child_ids" in vals:
                vals["child_ids"].append((0, 0, child_val))
            else:
                vals["child_ids"] = [(0, 0, child_val)]
        for i in LIST_COL_POP:
            if vals.get(i):
                vals.pop(i)
        return vals

    def partner_import_generic(self):
        fileobj = NamedTemporaryFile(
            "wb+", prefix="odoo-import_helper-", suffix=".xlsx"
        )
        file_bytes = base64.b64decode(self.file_import)
        fileobj.write(file_bytes)
        fileobj.seek(0)
        dataframe = opx.load_workbook(fileobj.name, read_only=True)
        reader = dataframe.active
        import_obj = self.env["import.helper"]
        speedy = import_obj._prepare_speedy(aiengine="NONE")
        line = 0
        colonnes = []
        partner_ids_list = self.speedy_partner_id()
        for row in reader.iter_rows(min_row=4, values_only=True):
            vals = {}
            if row[0] == "Colonnes:":
                for c in range(len(row)):
                    if row[c]:
                        colonnes.append(row[c])
                    else:
                        colonnes.append("empty")
                continue

            if row[2]:
                line += 1
                vals["line"] = line
                for c in range(len(row)):
                    if row[c] and colonnes[c] != "empty":
                        vals[colonnes[c]] = str(row[c])
            vals = self.check_vals_partner(vals)
            if vals["ref"] in partner_ids_list:
                import_obj._prepare_partner_vals(vals, speedy)
                record = self.env["res.partner"].browse(partner_ids_list[vals["ref"]])
                res = record.write(vals)
                if res:
                    logger.info(
                        f"{record.display_name},id {record.id} has been update with line {line}"
                    )
            else:
                import_obj._create_partner(vals, speedy)
        action = import_obj._result_action(speedy)
        return action

    def prepare_speedy_attribute_value(self):
        speedy_attribute_value = {}
        attribute_value = self.env["product.attribute.value"].search([])
        for att in attribute_value:
            speedy_attribute_value[att.fullname] = {
                "id": att.id,
                "attribute_id": att.attribute_id.id,
            }
        return speedy_attribute_value

    def check_vals_product(self, vals, speedy):
        variant_att = ()
        list_attribute_ids = {}
        template = False
        if "attributes" in vals:
            speedy_attribute_value = self.prepare_speedy_attribute_value()
            variant_att = vals["attributes"].split("/")
            for v in variant_att:
                if speedy_attribute_value.get(v):
                    if speedy_attribute_value[v]["attribute_id"] in list_attribute_ids:
                        list_attribute_ids[
                            speedy_attribute_value[v]["attribute_id"]
                        ].append(
                            speedy_attribute_value[v]["id"],
                        )
                    else:
                        list_attribute_ids[
                            speedy_attribute_value[v]["attribute_id"]
                        ] = [speedy_attribute_value[v]["id"]]
                else:
                    speedy["logs"]["product.product"].append(
                        {
                            "msg": f"Cannot found attrivutes {v} for {vals['line']}",
                            "value": v,
                            "vals": vals,
                            "field": "product.product,attribute_line_ids",
                        }
                    )

        if "ref_supplier" in vals:
            speedy_partner_id = self.speedy_partner_id()
            if vals["ref_supplier"] in speedy_partner_id:
                vals["supplier_id"] = speedy_partner_id[vals["ref_supplier"]]
        if vals.get("Colonnes:") == "product.template":
            template = True
        for i in LIST_COL_POP:
            if vals.get(i):
                vals.pop(i)
        return vals, variant_att, list_attribute_ids, template

    def product_seller_update(self, vals, record, speedy):
        supplierinfo_vals = {}
        if vals.get("supplier_id"):
            supplierinfo_vals = {
                "partner_id": vals["supplier_id"],
                "price": vals.get("supplier_price"),
                "product_code": vals.get("supplier_product_code"),
                "product_name": vals.get("supplier_product_name"),
                "min_qty": vals.get("supplier_min_qty"),
                "product_id": vals.get("supplier_product_id"),
            }
            if vals.get("supplier_delay"):
                supplierinfo_vals["delay"] = vals["supplier_delay"]
            if vals.get("supplier_currency"):
                if isinstance(vals["supplier_currency"], int):
                    supplierinfo_vals["currency_id"] = vals["supplier_currency"]
                elif isinstance(vals["supplier_currency"], str):
                    currency = vals["supplier_currency"].upper().strip()
                    if currency in speedy["currency2id"]:
                        supplierinfo_vals["currency_id"] = speedy["currency2id"][
                            currency
                        ]
                    else:
                        speedy["logs"]["product.product"].append(
                            {
                                "msg": "%s is not a known currency ISO code" % currency,
                                "value": currency,
                                "vals": vals,
                                "field": "product.supplierinfo,currency_id",
                                "reset": True,
                            }
                        )
        for seller in record.seller_ids:
            if seller.partner_id.id == supplierinfo_vals["partner_id"] and (
                seller.product_code == supplierinfo_vals.get("product_code")
                or seller.product_name == supplierinfo_vals.get("product_name")
            ):
                vals["seller_ids"] = [Command.update(seller.id, supplierinfo_vals)]
                vals.pop("supplier_id")
                return vals

        vals["seller_ids"] = [Command.create(supplierinfo_vals)]
        vals.pop("supplier_id")
        return vals

    def product_import_generic(self):
        fileobj = NamedTemporaryFile(
            "wb+", prefix="odoo-import_helper-", suffix=".xlsx"
        )
        file_bytes = base64.b64decode(self.file_import)
        fileobj.write(file_bytes)
        fileobj.seek(0)
        dataframe = opx.load_workbook(fileobj.name, read_only=True)
        reader = dataframe.active
        import_obj = self.env["import.helper"]
        speedy = import_obj._prepare_speedy(aiengine="NONE")
        line = 0
        colonnes = []
        reference = ""
        for row in reader.iter_rows(min_row=1, max_row=2, max_col=6, values_only=True):
            if row[0] == "Information":
                if row[1] == "Champ de reference":
                    reference = row[2]
        product_ids = self.env["product.product"].search_read([], [reference])
        speedy_product_list = {}
        for p in product_ids:
            speedy_product_list[p[reference]] = p["id"]
        product_template_ids = self.env["product.template"].search_read(
            [], ["default_code_import"]
        )
        speedy_product_template_list = {}
        for p in product_template_ids:
            speedy_product_template_list[p["default_code_import"]] = p["id"]
        list_product_create = {}
        count = 0
        for row in reader.iter_rows(min_row=4, values_only=True):
            vals = {}
            if count >= 200:
                self.env.cr.commit()
                count = 0
                logger.info("commit 200 product")
            if row[0] == "Colonnes:":
                for c in range(len(row)):
                    if row[c]:
                        colonnes.append(row[c])
                    else:
                        colonnes.append("empty")
                continue
            if row:
                line += 1
                count += 1
                vals["line"] = line
                for c in range(len(row)):
                    if row[c] and colonnes[c] != "empty":
                        vals[colonnes[c]] = str(row[c])
                if (
                    vals.get("ref_template")
                    and vals["ref_template"] in speedy_product_template_list
                ):
                    vals["product_tmpl_id"] = speedy_product_template_list[
                        vals["ref_template"]
                    ]
                vals, variant_att, list_attribue_ids, template = (
                    self.check_vals_product(vals, speedy)
                )
                # if (
                #     not vals.get("product_tmpl_id")
                #     and vals["default_code"] in speedy_product_template_list
                # ):
                #     location_id = vals.get("location_id") or speedy.get(
                #         "default_location_id"
                #     )
                #     vals = import_obj._prepare_product_vals(vals, location_id, speedy)
                #     res = (
                #         self.env["product.template"]
                #         .browse(speedy_product_list[vals["default_code"]])
                #         .write(vals)
                #     )
                #     if res:
                #         logger.info(
                #             f"{res.display_name},id {res.id} has been update with line {line}"
                #         )
                #     else:
                #         logger.warning(f"line {line} have done nothing")
                if not template:
                    if vals.get(reference) in speedy_product_list and self.update_on:
                        location_id = vals.get("location_id") or speedy.get(
                            "default_location_id"
                        )
                        ref_product = vals.pop(reference)
                        record = self.env["product.product"].browse(
                            speedy_product_list[ref_product]
                        )
                        if record:
                            if record.location_id and record.location_id != location_id:
                                location_id = record.location_id
                            if vals.get("default_code") == record.default_code:
                                vals.pop("default_code")
                            if vals.get("barcode") == record.barcode:
                                vals.pop("barcode")
                            if record.seller_ids:
                                if record.product_template_attribute_value_ids:
                                    vals["supplier_product_id"] = record.id
                                vals = self.product_seller_update(vals, record, speedy)
                            vals = import_obj._prepare_product_vals(
                                vals, location_id, speedy
                            )
                            if not vals:
                                logger.warning("Product on line %s skipped", line)
                                continue
                            if vals.get("standard_price"):
                                vals["standard_price"] = float(vals["standard_price"])
                            res = record.write(vals)
                            if res:
                                logger.info(
                                    f"{record.display_name}, id {record.id} has been update with line {line}"
                                )
                            else:
                                logger.warning(f"line {line} have done nothing")
                            continue
                    elif (
                        vals.get(reference) in speedy_product_list
                        and not self.update_on
                    ):
                        speedy["logs"]["product.product"].append(
                            {
                                "msg": f"Product with {vals[reference]} already exist for line {line}",
                                "value": vals[reference],
                                "vals": vals,
                                "field": f"product.product,{reference}",
                                "reset": True,
                            }
                        )
                    elif vals.get("product_tmpl_id"):
                        location_id = vals.get("location_id") or speedy.get(
                            "default_location_id"
                        )
                        if vals["product_tmpl_id"] in list_product_create:
                            template = list_product_create[vals["product_tmpl_id"]]
                        else:
                            template = self.env["product.template"].browse(
                                vals["product_tmpl_id"]
                            )
                        res_p = False
                        for p in template.product_variant_ids:
                            fullname_att = []
                            if p.product_template_attribute_value_ids:
                                for v in p.product_template_attribute_value_ids:
                                    fullname_att.append(
                                        v.product_attribute_value_id.fullname
                                    )

                                if fullname_att == variant_att and not p.barcode:
                                    vals["supplier_product_id"] = p.id
                                    vals = import_obj._prepare_product_vals(
                                        vals, location_id, speedy
                                    )
                                    if not vals:
                                        logger.warning(
                                            "Product on line %s skipped", line
                                        )
                                        continue
                                    if vals.get("standard_price"):
                                        vals["standard_price"] = float(
                                            vals["standard_price"]
                                        )
                                    if vals.get("list_price") and hasattr(
                                        p, "fix_price"
                                    ):
                                        vals["fix_price"] = vals["list_price"]
                                        vals.pop("list_price")
                                    elif vals.get("list_price"):
                                        vals.pop("list_price")
                                    res_p = p.write(vals)
                                    logger.info(
                                        f"product variant {p.id} has been update {fullname_att}"
                                    )
                                    break
                                elif p.barcode and fullname_att == variant_att:
                                    speedy["logs"]["product.product"].append(
                                        {
                                            "msg": f"{p.id} product with {variant_att} already exite",
                                            "value": variant_att,
                                            "vals": vals,
                                            "field": "product.product,attribute_line_ids",
                                            "reset": True,
                                        }
                                    )
                                    break
                        if not res_p:
                            speedy["logs"]["product.product"].append(
                                {
                                    "msg": f"Not product with {variant_att} for line {line}",
                                    "value": variant_att,
                                    "vals": vals,
                                    "field": "product.product,attribute_line_ids",
                                    "reset": True,
                                }
                            )

                        continue
                    else:
                        res = import_obj._create_product(vals, speedy)
                        continue
                elif template:
                    if vals.get(reference) in speedy_product_list and self.update_on:
                        location_id = vals.get("location_id") or speedy.get(
                            "default_location_id"
                        )
                        ref_product = vals.pop(reference)
                        record = self.env["product.product"].browse(
                            speedy_product_list[ref_product]
                        )
                        if record:
                            if record.location_id and record.location_id != location_id:
                                location_id = record.location_id
                            if vals.get("default_code") == record.default_code:
                                vals.pop("default_code")
                            if vals.get("barcode") == record.barcode:
                                vals.pop("barcode")
                            if record.seller_ids:
                                if record.product_template_attribute_value_ids:
                                    vals["supplier_product_id"] = record.id
                                vals = self.product_seller_update(vals, record, speedy)
                            vals = import_obj._prepare_product_vals(
                                vals, location_id, speedy
                            )
                            if not vals:
                                logger.warning("Product on line %s skipped", line)
                                continue
                            if vals.get("standard_price"):
                                vals["standard_price"] = float(vals["standard_price"])
                            res = record.write(vals)
                            if res:
                                logger.info(
                                    f"{record.display_name}, id {record.id} has been update with line {line}"
                                )
                            else:
                                logger.warning(f"line {line} have done nothing")
                            continue
                    elif (
                        vals.get(reference) in speedy_product_list
                        and not self.update_on
                    ):
                        speedy["logs"]["product.product"].append(
                            {
                                "msg": f"Product with {vals[reference]} already exist for line {line}",
                                "value": vals[reference],
                                "vals": vals,
                                "field": f"product.product,{reference}",
                                "reset": True,
                            }
                        )
                        continue
                    elif (
                        vals.get("default_code") in speedy_product_template_list
                        and self.update_on
                    ):
                        location_id = vals.get("location_id") or speedy.get(
                            "default_location_id"
                        )
                        record = self.env["product.template"].browse(
                            speedy_product_template_list[vals["default_code"]]
                        )
                        ref_product = vals.pop("default_code")
                        if record:
                            if record.location_id and record.location_id != location_id:
                                location_id = record.location_id
                            if vals.get("barcode") == record.barcode:
                                vals.pop("barcode")
                            if record.seller_ids:
                                vals = self.product_seller_update(vals, record, speedy)
                            vals = import_obj._prepare_product_vals(
                                vals, location_id, speedy
                            )
                            if not vals:
                                logger.warning("Product on line %s skipped", line)
                                continue

                            res = record.write(vals)
                            if res:
                                if variant_att:
                                    speedy_line_attr = {}
                                    for line_id in record.attribute_line_ids:
                                        speedy_line_attr[line_id.attribute_id.id] = {
                                            "id": line_id.id,
                                            "value_ids": line_id.value_ids.ids,
                                        }

                                    for att in list_attribue_ids:
                                        b = list_attribue_ids[att]
                                        if (
                                            speedy_line_attr.get(att)
                                            and b != speedy_line_attr[att]["value_ids"]
                                        ):
                                            record.attribute_line_ids = [
                                                Command.update(
                                                    speedy_line_attr[att]["id"],
                                                    {
                                                        "value_ids": [
                                                            Command.set(
                                                                speedy_line_attr[att][
                                                                    "value_ids"
                                                                ]
                                                            )
                                                        ]
                                                    },
                                                )
                                            ]
                                        elif (
                                            speedy_line_attr.get(att)
                                            and b == speedy_line_attr[att]["value_ids"]
                                        ):
                                            continue
                                        else:
                                            record.attribute_line_ids = [
                                                Command.create(
                                                    {
                                                        "attribute_id": att,
                                                        "value_ids": [Command.set(b)],
                                                    }
                                                )
                                            ]
                                    record.default_code = ref_product
                                logger.info(f"Update {record.name} {record.id} Ok")
                                continue
                            else:
                                logger.warning(
                                    f"ERREUR lors de la mise a jour du product line {line}"
                                )
                                continue
                        else:
                            logger.warning(f"No product found for {line}")
                    elif (
                        vals.get("default_code") in speedy_product_template_list
                        and not self.update_on
                    ):
                        speedy["logs"]["product.product"].append(
                            {
                                "msg": f"Product_Template with {vals['default_code']} already exist for line {line}",
                                "value": vals["default_code"],
                                "vals": vals,
                                "field": "product.product,default_code",
                                "reset": True,
                            }
                        )
                    else:
                        location_id = vals.get("location_id") or speedy.get(
                            "default_location_id"
                        )
                        vals = import_obj._prepare_product_vals(
                            vals, location_id, speedy
                        )
                        if not vals:
                            logger.warning("Product on line %s skipped", line)
                            continue
                        p_tmpl = self.env["product.template"].create(vals)
                        speedy_product_template_list[p_tmpl.default_code] = p_tmpl.id
                        list_product_create[p_tmpl.id] = p_tmpl
                        if p_tmpl and variant_att:
                            for att in list_attribue_ids:
                                b = list_attribue_ids[att]
                                p_tmpl.attribute_line_ids = [
                                    Command.create(
                                        {
                                            "attribute_id": att,
                                            "value_ids": [Command.set(b)],
                                        }
                                    )
                                ]
                                logger.info(f"product variant for {att} create")
                            p_tmpl.default_code = vals["default_code"]
                            logger.info(
                                f"{p_tmpl.id} has been create with {len(p_tmpl.attribute_line_ids)} variant"
                            )
                        else:
                            logger.info(f"{p_tmpl.id} has been create")
                            continue
                # elif (not template or template) and not variant_att:
                #     res = import_obj._create_product(vals, speedy)
                #     continue
                else:
                    logger.warning(f"NO PRODUCT IMPORTED line {line} Name {row[1]}")
                    speedy["logs"]["product.product"].append(
                        {
                            "msg": "Pas d'identification product ou template colonnes type [A]",
                            "value": row[0],
                            "vals": vals,
                            "field": "product.product,product_tmpl_id",
                            "reset": True,
                        }
                    )
            else:
                break
        for t in speedy_product_template_list:
            record = self.env["product.template"].browse(
                speedy_product_template_list[t]
            )
            record.default_code = t
        action = import_obj._result_action(speedy)
        return action
